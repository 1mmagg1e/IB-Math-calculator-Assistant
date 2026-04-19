import os
import urllib.parse
from docx import Document  # 导入操作 Word 的库


# ==================================
### in terminal enter "pip install python-docx"
# 1. 你本地存放图片的文件夹名字
image_folder = 'image'

# 2. GitHub Raw 链接的前缀（注意最后要保留斜杠 /）
github_base_url = 'https://raw.githubusercontent.com/1mmagg1e/ti84_calc_image/refs/heads/main/images/'

# 3. 最后生成的文本文件名 (改成 .docx)
output_file = 'Image_Index.docx'
# =========================================

# 检查图片文件夹是否存在
if not os.path.exists(image_folder):
    print(f"❌ fail find file：{image_folder}，please check route！")
else:
    # 创建一个空白的 Word 文档
    doc = Document()

    # 写入大标题 (Heading 1)
    doc.add_heading('TI-84 visual search dictionary', level=1)

    count = 0
    # 遍历图片文件夹里的所有文件
    for filename in os.listdir(image_folder):
        # 只处理图片文件
        if filename.lower().endswith(('.png', '.jpg', '.jpeg')):
            count += 1

            # 提取按键名字
            button_name = os.path.splitext(filename)[0]

            # 安全处理文件名
            safe_filename = urllib.parse.quote(filename)
            full_url = f"{github_base_url}{safe_filename}"

            # 把按键名字作为三级标题写进 Word
            doc.add_heading(button_name, level=3)

            # 把图片链接作为普通段落写进 Word
            markdown_line = f"![{button_name}]({full_url})"
            doc.add_paragraph(markdown_line)

            # 增加一个空段落，这在 Dify 里会被精准识别为 \n\n
            doc.add_paragraph("")

    # 保存文档
    doc.save(output_file)

    print(f"✅ Finished! Successfully generated Word Document for {count} image files.")
    print(f"📄 Please check the {output_file} file.")
